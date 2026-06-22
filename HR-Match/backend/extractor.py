import fitz
from docx import Document
from pathlib import Path


def extract_pdf(file_path):
    result = {
        "filename": "",
        "text": "",
        "page_count": 0,
        "type": "pdf",
        "success": False,
        "error": None
    }
    try:
        doc = fitz.open(file_path)
        result["filename"] = Path(file_path).name
        full_text = []
        for page in doc:
            full_text.append(page.get_text())
        result["page_count"] = len(doc)
        result["text"] = "\n---\n".join(full_text)
        result["success"] = True
        doc.close()
    except Exception as e:
        result["error"] = str(e)
    return result


def extract_docx(file_path):
    result = {
        "filename": "",
        "text": "",
        "paragraph_count": 0,
        "type": "docx",
        "success": False,
        "error": None
    }
    try:
        doc = Document(file_path)
        result["filename"] = Path(file_path).name
        paragraphs = []
        for p in doc.paragraphs:
            line = p.text.strip()
            if line:
                paragraphs.append(line)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(c for c in cells if c)
                if row_text:
                    paragraphs.append(row_text)
        result["paragraph_count"] = len(paragraphs)
        result["text"] = "\n".join(paragraphs)
        result["success"] = True
    except Exception as e:
        result["error"] = str(e)
    return result


def extract_file(file_path):
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(file_path)
    elif suffix == ".docx":
        return extract_docx(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {suffix}")


def extract_folder(folder_path):
    results = []
    folder = Path(folder_path)
    for f in folder.iterdir():
        if f.suffix.lower() in (".pdf", ".docx"):
            r = extract_file(str(f))
            results.append(r)
    results.sort(key=lambda r: r["filename"])
    return results
